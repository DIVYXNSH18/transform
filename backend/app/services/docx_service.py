import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def _parse_markdown_to_docx(doc, markdown_text: str):
    """
    Parses a markdown string and renders it into a python-docx Document,
    handling headings (#, ##, ###), bold (**text**), bullet points (- or •),
    and markdown tables (| col | col |).
    """
    lines = markdown_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Skip empty lines
        if not stripped:
            i += 1
            continue

        # Headings
        if stripped.startswith("### "):
            h = doc.add_heading(level=3)
            _add_bold_runs(h, stripped[4:].strip())
            i += 1
            continue
        if stripped.startswith("## "):
            h = doc.add_heading(level=2)
            _add_bold_runs(h, stripped[3:].strip())
            i += 1
            continue
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            _add_bold_runs(h, stripped[2:].strip())
            i += 1
            continue

        # Markdown table detection
        if "|" in stripped and stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            _render_table(doc, table_lines)
            continue

        # Bullet points
        if stripped.startswith("- ") or stripped.startswith("• ") or stripped.startswith("* "):
            text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            _add_bold_runs(p, text)
            i += 1
            continue

        # Blockquotes / Executive Callouts
        if stripped.startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip("> ").strip())
                i += 1
            quote_text = " ".join(quote_lines)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            _add_bold_runs(p, quote_text)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(15, 23, 42)
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        _add_bold_runs(p, stripped)
        i += 1

def _add_bold_runs(paragraph, text: str):
    """Splits text on **bold** markers and adds runs with appropriate formatting."""
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(10)
        else:
            run = paragraph.add_run(part)
            run.font.name = "Arial"
            run.font.size = Pt(10)

def _render_table(doc, table_lines: list):
    """Renders a markdown table (| col | col |) into a Word table."""
    # Filter out separator rows (|---|---|)
    data_rows = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip("|").split("|")]
        if cells and not all(re.match(r'^[\-:]+$', c) for c in cells):
            data_rows.append(cells)

    if not data_rows:
        return

    num_cols = len(data_rows[0])
    table = doc.add_table(rows=len(data_rows), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    for row_idx, row_data in enumerate(data_rows):
        for col_idx, cell_text in enumerate(row_data):
            if col_idx < num_cols:
                cell = table.rows[row_idx].cells[col_idx]
                cell.text = cell_text.strip()
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(9)
                        if row_idx == 0:
                            run.bold = True

def create_executive_docx(title: str, summary_markdown: str, ico_data: dict, output_path: str):
    """
    Builds a professional executive document (.docx) using python-docx.
    Renders the full LLM-generated summary_markdown into the document,
    with ICO metadata as a header section.
    """
    doc = Document()
    
    # Configure 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Header
    p_meta = doc.add_paragraph()
    run_meta = p_meta.add_run("TRANSFORMAI // EXECUTIVE DELIVERABLE BRIEF")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(9)
    run_meta.font.bold = True
    run_meta.font.color.rgb = RGBColor(255, 107, 0) # Performance Orange

    # Main Title
    h1 = doc.add_heading(level=1)
    run_h1 = h1.add_run(title)
    run_h1.font.name = "Arial"
    run_h1.font.size = Pt(22)
    run_h1.font.bold = True
    run_h1.font.color.rgb = RGBColor(15, 23, 42)

    # Sub-metadata bar
    p_info = doc.add_paragraph()
    timestamp = ico_data.get("timestamp", "Present")
    loc = ico_data.get("location", "Edge Engine")
    obj = ico_data.get("primary_objective", "Operational alignment")
    p_info.add_run(f"Timestamp: {timestamp}  |  Location: {loc}\nObjective: {obj}\n").italic = True

    # Render full LLM-generated summary if available
    if summary_markdown and summary_markdown.strip():
        doc.add_paragraph()  # spacer
        _parse_markdown_to_docx(doc, summary_markdown)
    else:
        # Fallback: render from ICO fields directly
        doc.add_heading("1. Strategic Overview", level=2)
        overview_text = ico_data.get("executive_overview", "")
        p_over = doc.add_paragraph(overview_text)
        p_over.style.font.name = "Arial"

        doc.add_heading("2. Key Findings & Observations", level=2)
        findings = ico_data.get("key_findings", [])
        for f in findings:
            doc.add_paragraph(f, style='List Bullet')

        doc.add_heading("3. Action Matrix", level=2)
        actions = ico_data.get("action_items", [])
        if actions:
            table = doc.add_table(rows=1, cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Owner"
            hdr_cells[1].text = "Action Item"
            hdr_cells[2].text = "Deadline"

            for item in actions:
                row_cells = table.add_row().cells
                row_cells[0].text = item.get("owner", "Team")
                row_cells[1].text = item.get("task", "")
                row_cells[2].text = item.get("deadline", "TBD")

    # Entities / Metrics
    entities = ico_data.get("entities", {})
    if isinstance(entities, dict):
        metrics = entities.get("metrics", [])
    else:
        metrics = []
    if metrics:
        doc.add_heading("Key Milestones & Metrics", level=2)
        for m in metrics:
            doc.add_paragraph(f"Target: {m}", style='List Bullet')

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    return output_path
